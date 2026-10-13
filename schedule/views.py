import calendar
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.contrib import messages
from .models import WorkSchedule, ScheduleEntry, SHIFT_HOURS

@login_required
def schedule_list(request):
    schedules = WorkSchedule.objects.all().order_by('-year','-month')
    return render(request, 'schedule/list.html', {'schedules': schedules})

@login_required
def schedule_create(request):
    if request.method == 'POST':
        brigade_no = request.POST.get('brigade_no') or 4
        month = int(request.POST.get('month') or 10)
        year = int(request.POST.get('year') or 2026)
        norm_hours = float(str(request.POST.get('norm_hours') or 176).replace(',','.'))
        sched = WorkSchedule.objects.create(brigade_no=brigade_no, month=month, year=year, norm_hours=norm_hours)
        emp_ids = request.POST.getlist('employees')
        days = calendar.monthrange(year, month)[1]
        bulk = []
        for emp_id in emp_ids:
            for d in range(1, days+1):
                bulk.append(ScheduleEntry(schedule=sched, employee_id=emp_id, day=d, shift='', hours=0))
        ScheduleEntry.objects.bulk_create(bulk)
        return redirect('schedule_edit', pk=sched.id)
    employees = User.objects.all().order_by('last_name','first_name')
    return render(request, 'schedule/form.html', {'employees': employees})

@login_required
def schedule_delete(request, pk):
    sched = get_object_or_404(WorkSchedule, pk=pk)
    if request.method == 'POST':
        sched.delete()
        messages.success(request, 'График удален')
        return redirect('schedule_list')
    return redirect('schedule_list')

@login_required
def schedule_add_employee(request, pk):
    sched = get_object_or_404(WorkSchedule, pk=pk)
    if request.method == 'POST':
        emp_id = request.POST.get('employee_id')
        if emp_id and not ScheduleEntry.objects.filter(schedule=sched, employee_id=emp_id).exists():
            days = calendar.monthrange(sched.year, sched.month)[1]
            bulk = [ScheduleEntry(schedule=sched, employee_id=emp_id, day=d, shift='', hours=0) for d in range(1, days+1)]
            ScheduleEntry.objects.bulk_create(bulk)
            messages.success(request, 'Сотрудник добавлен')
    return redirect('schedule_edit', pk=pk)

@login_required
def schedule_remove_employee(request, pk, user_id):
    sched = get_object_or_404(WorkSchedule, pk=pk)
    if request.method == 'POST':
        ScheduleEntry.objects.filter(schedule=sched, employee_id=user_id).delete()
        messages.success(request, 'Сотрудник удален из графика')
    return redirect('schedule_edit', pk=pk)

@login_required
def schedule_edit(request, pk):
    sched = get_object_or_404(WorkSchedule, pk=pk)
    days = list(range(1, calendar.monthrange(sched.year, sched.month)[1]+1))
    employees = User.objects.filter(scheduleentry__schedule=sched).distinct().order_by('last_name','first_name')
    if request.method == 'POST' and any(k.startswith('shift_') for k in request.POST.keys()):
        for emp in employees:
            for d in days:
                shift = request.POST.get(f'shift_{emp.id}_{d}','').strip().upper()
                hours_raw = request.POST.get(f'hours_{emp.id}_{d}','0').replace(',','.')
                try: hours = float(hours_raw) if hours_raw else 0
                except: hours = 0
                if hours==0 and shift in ('Д','П','У','Н'): hours = SHIFT_HOURS.get(shift,0)
                if shift in ('В','О',''): hours = 0
                obj, _ = ScheduleEntry.objects.get_or_create(schedule=sched, employee=emp, day=d)
                obj.shift = shift; obj.hours = hours; obj.save()
        messages.success(request, 'Сохранено')
        return redirect('schedule_edit', pk=pk)
    entries = ScheduleEntry.objects.filter(schedule=sched)
    matrix = {}
    totals = {}
    for e in entries: matrix.setdefault(e.employee_id, {})[e.day] = e
    norm = float(sched.norm_hours or 176)
    for emp in employees:
        total = sum([float(x.hours or 0) for x in entries.filter(employee=emp)])
        ur = min(total, norm); sv = max(0, total - norm)
        totals[emp.id] = {'ur': f"{ur:.1f}".replace('.',','), 'sv': f"{sv:.1f}".replace('.',',') if sv else "0"}
    existing_ids = employees.values_list('id', flat=True)
    available_employees = User.objects.exclude(id__in=existing_ids).order_by('last_name','first_name')
    return render(request, 'schedule/edit.html', {'schedule': sched,'employees': employees,'days': days,'matrix': matrix,'totals': totals,'available_employees': available_employees})

@login_required
def schedule_excel(request, pk):
    import calendar
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from django.http import HttpResponse
    sched = get_object_or_404(WorkSchedule, pk=pk)
    days_count = calendar.monthrange(sched.year, sched.month)[1]
    days = list(range(1, days_count+1))
    employees = User.objects.filter(scheduleentry__schedule=sched).distinct().order_by('last_name','first_name')
    entries = ScheduleEntry.objects.filter(schedule=sched)
    matrix = {}
    for e in entries:
        matrix.setdefault(e.employee_id, {})[e.day] = e

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Бригада {sched.brigade_no} {sched.month}-{sched.year}"

    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=days_count+3)
    ws.cell(row=1, column=1, value=f"График работы Бригады №{sched.brigade_no} на {sched.month} / {sched.year} года. Норма {sched.norm_hours}ч").font = Font(bold=True, size=12)

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=days_count+3)
    ws.cell(row=2, column=1, value="Утверждено: ____________  Д-07:00-15:45, П-15:45-00:15, У-07:00-19:00, Н-19:00-07:00, В-выходной, О-отпуск").font = Font(size=9, italic=True)

    header = ["ФИО"] + [str(d) for d in days] + ["Ур.", "Св.ур."]
    header_row = 4
    thin = Side(style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill_header = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
    fill_gray = PatternFill(start_color="D1D5DB", end_color="D1D5DB", fill_type="solid")

    for col_idx, val in enumerate(header, start=1):
        c = ws.cell(row=header_row, column=col_idx, value=val)
        c.font = Font(bold=True, size=10)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.fill = fill_header
        c.border = border

    # Column widths
    ws.column_dimensions[openpyxl.utils.get_column_letter(1)].width = 28
    for i in range(2, days_count+2):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 4.5
    ws.column_dimensions[openpyxl.utils.get_column_letter(days_count+2)].width = 7
    ws.column_dimensions[openpyxl.utils.get_column_letter(days_count+3)].width = 7

    row = header_row + 1
    norm = float(sched.norm_hours or 176)
    for emp in employees:
        # ФИО
        try:
            fio = emp.profile.short_name if hasattr(emp, 'profile') and emp.profile.short_name else f"{emp.last_name} {emp.first_name}".strip() or emp.username
        except:
            fio = emp.username

        total = sum([float(matrix.get(emp.id, {}).get(d).hours or 0) if matrix.get(emp.id, {}).get(d) else 0 for d in days])
        ur = min(total, norm)
        sv = max(0, total-norm)

        # Two rows per employee: row shift, row+1 hours
        ws.merge_cells(start_row=row, start_column=1, end_row=row+1, end_column=1)
        fio_cell = ws.cell(row=row, column=1, value=fio)
        fio_cell.font = Font(bold=True, size=10)
        fio_cell.alignment = Alignment(horizontal='left', vertical='center')
        fio_cell.border = border

        # second cell for merged area needs border too
        ws.cell(row=row+1, column=1).border = border

        for d in days:
            entry = matrix.get(emp.id, {}).get(d)
            shift = entry.shift if entry else ""
            hours = entry.hours if entry and entry.hours else 0

            c1 = ws.cell(row=row, column=1+d, value=shift)
            c1.alignment = Alignment(horizontal='center', vertical='center')
            c1.font = Font(bold=True, size=10)
            c1.border = border
            if shift in ('В','О'):
                c1.fill = fill_gray

            c2 = ws.cell(row=row+1, column=1+d, value=float(hours) if hours else "")
            c2.alignment = Alignment(horizontal='center', vertical='center')
            c2.font = Font(size=9)
            c2.border = border
            c2.number_format = '0.0'
            if shift in ('В','О'):
                c2.fill = fill_gray

        # Totals - merged vertically
        ws.merge_cells(start_row=row, start_column=days_count+2, end_row=row+1, end_column=days_count+2)
        ws.merge_cells(start_row=row, start_column=days_count+3, end_row=row+1, end_column=days_count+3)
        ur_cell = ws.cell(row=row, column=days_count+2, value=round(ur,1))
        ur_cell.font = Font(bold=True)
        ur_cell.alignment = Alignment(horizontal='center', vertical='center')
        ur_cell.border = border
        sv_cell = ws.cell(row=row, column=days_count+3, value=round(sv,1) if sv else 0)
        sv_cell.font = Font(bold=True)
        sv_cell.alignment = Alignment(horizontal='center', vertical='center')
        sv_cell.border = border
        ws.cell(row=row+1, column=days_count+2).border = border
        ws.cell(row=row+1, column=days_count+3).border = border

        row += 2

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=brigada_{sched.brigade_no}_{sched.month}_{sched.year}.xlsx'
    wb.save(response)
    return response
